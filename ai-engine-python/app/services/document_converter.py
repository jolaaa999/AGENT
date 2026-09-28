"""
文档转换为 Markdown

支持把常见学习资料格式统一转成 Markdown 笔记，作为图谱生成的输入：
- .md / .txt          → 直接读取（自动尝试多种编码）
- .docx               → python-docx 提取段落与表格
- .pdf                → pypdf 逐页提取文本（扫描版 PDF 提取不到文字，会给出提示）
- 图片（png/jpg/...）  → 交给多模态 LLM 识图，抽取关键信息为 Markdown

设计原则：
- 纯本地解析优先（不消耗额度、速度快），只有图片才调用 LLM。
- 解析失败时抛出带有可读原因的 DocumentConvertError，由上层转成 400，
  而不是让 500 掩盖「这个文件确实读不了」的事实。
"""

from __future__ import annotations

import base64
import io
import logging
import os
from typing import Optional

logger = logging.getLogger(__name__)


class DocumentConvertError(Exception):
    """文档无法转换时抛出，message 面向终端用户。"""


# 图片扩展名 → MIME 类型
IMAGE_MIME_TYPES = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
    ".gif": "image/gif",
    ".bmp": "image/bmp",
}

# 单文件大小上限（与前端提示保持一致，防止超大文件打爆内存）
MAX_FILE_BYTES = 20 * 1024 * 1024

# PDF 提取文本少于该字符数时，判定为可能是扫描件
PDF_MIN_TEXT_CHARS = 20


def _decode_text(raw: bytes) -> str:
    """按常见编码依次尝试解码文本文件，避免中文笔记因 GBK 而乱码。"""
    for enc in ("utf-8-sig", "utf-8", "gb18030", "gbk", "big5", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    raise DocumentConvertError("无法识别文件编码，请另存为 UTF-8 后重试")


def _convert_docx(raw: bytes) -> str:
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover
        raise DocumentConvertError("服务器未安装 python-docx，无法解析 .docx") from exc

    try:
        doc = docx.Document(io.BytesIO(raw))
    except Exception as exc:
        raise DocumentConvertError(
            "无法打开该 Word 文档。若为 .doc（旧版二进制格式），请先另存为 .docx"
        ) from exc

    lines: list[str] = []
    for para in doc.paragraphs:
        text = (para.text or "").strip()
        if not text:
            continue
        # 保留标题层级，让后续 AST 分块能识别结构
        style = (para.style.name or "") if para.style is not None else ""
        if style.startswith("Heading"):
            level = "".join(ch for ch in style if ch.isdigit())
            prefix = "#" * min(int(level or 1), 6)
            lines.append(f"{prefix} {text}")
        else:
            lines.append(text)

    # 表格转成 Markdown 表格，避免丢失结构化内容
    for table in doc.tables:
        rows = [[(cell.text or "").strip().replace("\n", " ") for cell in row.cells] for row in table.rows]
        if not rows or not any(any(c for c in row) for row in rows):
            continue
        width = max(len(r) for r in rows)
        lines.append("")
        for idx, row in enumerate(rows):
            padded = row + [""] * (width - len(row))
            lines.append("| " + " | ".join(padded) + " |")
            if idx == 0:
                lines.append("| " + " | ".join(["---"] * width) + " |")
        lines.append("")

    return "\n".join(lines).strip()


def _convert_pdf(raw: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise DocumentConvertError("服务器未安装 pypdf，无法解析 PDF") from exc

    try:
        reader = PdfReader(io.BytesIO(raw))
    except Exception as exc:
        raise DocumentConvertError("无法打开该 PDF（文件可能已损坏或加密）") from exc

    if getattr(reader, "is_encrypted", False):
        # 很多 PDF 只是带「加密标记」而没有真实口令（打印/复制限制等），
        # 用空口令尝试解密是合理的；但必须检查返回值。
        #
        # pypdf 的 decrypt() 在口令错误时返回 0（PasswordType.NOT_DECRYPTED），
        # **不会抛异常**。原实现忽略了返回值，于是带真实口令的 PDF 会继续走到
        # reader.pages，抛出未捕获的 FileNotDecryptedError，被上层
        # `except Exception` 包装成「文档转换失败: File has not been decrypted」
        # ——用户看到的是英文内部信息，而不是可读的中文提示。
        try:
            result = reader.decrypt("")
        except Exception as exc:
            raise DocumentConvertError("该 PDF 已加密，请先解除密码保护") from exc
        # pypdf 6.x 实测：解密成功后 is_encrypted 仍为 True，
        # 因此**只能**依据 decrypt() 的返回值判断是否真的解密成功。
        # PasswordType.NOT_DECRYPTED == 0 表示失败；其余（USER/OWNER_PASSWORD）为成功。
        # 注意 result 是 IntEnum，不能用 `not result` 之外的宽松判断，
        # 也不能再去看 is_encrypted。
        if not int(result):
            raise DocumentConvertError(
                "该 PDF 已加密（需要密码），请先解除密码保护，"
                "或另存为未加密版本后重试"
            )

    pages: list[str] = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            text = (page.extract_text() or "").strip()
        except Exception:
            text = ""
        if text:
            pages.append(text)

    content = "\n\n".join(pages).strip()
    if len(content) < PDF_MIN_TEXT_CHARS:
        raise DocumentConvertError(
            "该 PDF 提取不到文字，可能是扫描件/图片型 PDF。"
            "请截取页面另存为图片后上传，我会用识图方式提取内容"
        )
    return content


def _convert_image(raw: bytes, filename: str) -> str:
    """图片走多模态 LLM 识图，抽取关键信息为 Markdown。"""
    from app.core.llm import build_chat_openai
    from langchain_core.messages import HumanMessage, SystemMessage

    ext = os.path.splitext(filename)[1].lower()
    mime = IMAGE_MIME_TYPES.get(ext, "image/png")
    b64 = base64.b64encode(raw).decode()

    llm = build_chat_openai(temperature=0)
    prompt = (
        "请识别这张图片中的学习资料内容，并整理成结构清晰的 Markdown 笔记。\n"
        "要求：\n"
        "1. 忠实还原图中的文字、公式、表格与结构，不要编造图中没有的内容\n"
        "2. 使用标题层级组织内容，公式用 LaTeX（行内 $...$，块级 $$...$$）\n"
        "3. 表格用 Markdown 表格表示，代码/伪代码用代码块\n"
        "4. 只输出 Markdown 正文，不要额外的开场白或解释"
    )
    message = HumanMessage(
        content=[
            {"type": "text", "text": prompt},
            {"type": "image_url", "image_url": {"url": f"data:{mime};base64,{b64}"}},
        ]
    )
    try:
        resp = llm.invoke([SystemMessage(content="你是专业的学习资料数字化助手。"), message])
    except Exception as exc:
        raise DocumentConvertError(f"识图失败：{exc}") from exc

    content = getattr(resp, "content", "") or ""
    if isinstance(content, list):
        content = "".join(
            part.get("text", "") if isinstance(part, dict) else str(part) for part in content
        )
    content = content.strip()
    if not content:
        raise DocumentConvertError("识图没有返回可用内容，请换一张更清晰的图片重试")
    return content


def convert_to_markdown(filename: str, raw: bytes) -> tuple[str, str]:
    """
    把上传文件转换为 Markdown。

    Returns:
        (markdown, 来源说明) —— 来源说明用于前端提示，例如 "docx"、"pdf"、"image"、"text"

    Raises:
        DocumentConvertError: 格式不支持或解析失败（message 面向用户）
    """
    if not raw:
        raise DocumentConvertError("文件内容为空")
    if len(raw) > MAX_FILE_BYTES:
        raise DocumentConvertError(f"文件过大（>{MAX_FILE_BYTES // 1024 // 1024}MB），请拆分后上传")

    ext = os.path.splitext(filename)[1].lower()

    if ext in IMAGE_MIME_TYPES:
        return _convert_image(raw, filename), "image"
    if ext == ".docx":
        return _convert_docx(raw), "docx"
    if ext == ".doc":
        raise DocumentConvertError("暂不支持旧版 .doc 格式，请用 Word 另存为 .docx 后上传")
    if ext == ".pdf":
        return _convert_pdf(raw), "pdf"
    if ext in (".md", ".markdown", ".txt", ""):
        return _decode_text(raw), "text"

    raise DocumentConvertError(
        f"暂不支持 {ext or '该'} 格式。支持：.md / .txt / .docx / .pdf / 图片（png、jpg、jpeg、webp）"
    )
